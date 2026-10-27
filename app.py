from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any

from pypdf import PdfReader
from docx import Document
from flask import Flask, flash, redirect, render_template, request, url_for
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.utils import secure_filename

from matcher import analyze_candidates

BASE_DIR = Path(__file__).resolve().parent
ALLOWED_EXTENSIONS = {"pdf", "docx"}
MAX_RESUMES_PER_RUN = 10

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("RESUME_APP_SECRET", "local-demo-key-change-before-deployment")
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16 MB total request size


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract selectable text from a PDF held in memory; do not save the upload."""
    chunks: list[str] = []
    reader = PdfReader(io.BytesIO(file_bytes))
    for page in reader.pages:
        chunks.append(page.extract_text() or "")
    return "\n".join(chunks).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraph and table text from a DOCX held in memory."""
    document = Document(io.BytesIO(file_bytes))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cell_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_text:
                chunks.append(" | ".join(cell_text))
    return "\n".join(chunks).strip()


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    extension = filename.rsplit(".", 1)[1].lower()
    if extension == "pdf":
        return extract_pdf_text(file_bytes)
    if extension == "docx":
        return extract_docx_text(file_bytes)
    raise ValueError("Unsupported file type. Upload a PDF or DOCX file.")


@app.get("/")
def index():
    return render_template("index.html", results=None, summary=None)


@app.post("/analyze")
def analyze():
    job_title = request.form.get("job_title", "").strip()
    job_description = request.form.get("job_description", "").strip()
    required_skills = request.form.get("required_skills", "").strip()
    uploaded_files = request.files.getlist("resumes")

    if not job_title:
        flash("Please enter a job title.", "error")
        return redirect(url_for("index"))
    if not job_description:
        flash("Please enter the job description.", "error")
        return redirect(url_for("index"))

    usable_files = [file for file in uploaded_files if file and file.filename]
    if not usable_files:
        flash("Upload at least one PDF or DOCX resume to run the analysis.", "error")
        return redirect(url_for("index"))
    if len(usable_files) > MAX_RESUMES_PER_RUN:
        flash(f"Please upload no more than {MAX_RESUMES_PER_RUN} resumes per analysis.", "error")
        return redirect(url_for("index"))

    candidates: list[dict[str, Any]] = []
    file_errors: list[str] = []

    for upload in usable_files:
        original_name = upload.filename or "candidate_resume"
        safe_name = secure_filename(original_name) or "candidate_resume"
        if not allowed_file(safe_name):
            file_errors.append(f"{original_name}: only PDF and DOCX files are supported.")
            continue
        try:
            file_bytes = upload.read()
            if not file_bytes:
                file_errors.append(f"{original_name}: the file is empty.")
                continue
            resume_text = extract_resume_text(safe_name, file_bytes)
            if len(resume_text.strip()) < 30:
                file_errors.append(
                    f"{original_name}: not enough text could be extracted. If it is a scanned PDF, use a text-based PDF or DOCX."
                )
                continue
            display_name = Path(safe_name).stem.replace("_", " ").replace("-", " ").strip() or "Candidate"
            candidates.append({"filename": safe_name, "name": display_name, "text": resume_text})
        except Exception as exc:  # Keep one malformed resume from stopping the whole batch.
            app.logger.warning("Could not process uploaded resume %s: %s", safe_name, exc)
            file_errors.append(f"{original_name}: the file could not be read. Try opening it and saving it again.")

    if not candidates:
        message = "No valid resumes could be analyzed. Check the file types and try again."
        if file_errors:
            message += " " + " ".join(file_errors)
        flash(message, "error")
        return redirect(url_for("index"))

    job_text = f"{job_title}\n{job_description}"
    results = analyze_candidates(job_text, candidates, required_skills)
    results.sort(key=lambda item: item["match_score"], reverse=True)
    for rank, result in enumerate(results, start=1):
        result["rank"] = rank

    summary = {
        "job_title": job_title,
        "analyzed": len(results),
        "with_errors": len(file_errors),
        "average_score": round(sum(item["match_score"] for item in results) / len(results)),
        "shortlisted": sum(1 for item in results if item["match_score"] >= 50),
        "scoring_note": (
            "Score = 65% text similarity + 35% required-skill coverage when required skills are entered. "
            "Without a required-skill list, the score uses text similarity only."
        ),
    }
    return render_template(
        "index.html",
        results=results,
        summary=summary,
        submitted_job_title=job_title,
        submitted_job_description=job_description,
        submitted_required_skills=required_skills,
        file_errors=file_errors,
    )


@app.errorhandler(RequestEntityTooLarge)
def handle_too_large(_error):
    flash("The total upload is too large. Keep the total request below 16 MB.", "error")
    return redirect(url_for("index")), 413


if __name__ == "__main__":
    # Local development only. Do not use Flask's debug server for public deployment.
    app.run(host="127.0.0.1", port=5000, debug=True)
